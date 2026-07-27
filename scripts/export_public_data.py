"""Export public downloads for the Mapping Viz gallery."""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SITE_ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs"
    / "019f8b8b-1bb8-73c2-a6d9-ab5b8ac2baea"
    / "Mapping_v2_analysis_ready.xlsx"
)
PROJECT_OUTPUT = SITE_ROOT / "data" / "mapping-bdph-projects.csv"
AGGREGATE_SOURCE = SITE_ROOT / "data" / "mapping-visualizations.js"
AGGREGATE_OUTPUT = SITE_ROOT / "data" / "mapping-visualization-aggregates.json"

PROJECT_FIELDS = [
    ("BDPH ID", "Project ID"),
    ("Name", "Project Name"),
    ("Project URL", "Project URL"),
    ("About/Summary", "Summary"),
    ("Project Start (Raw)", "Project Start"),
    ("Start Year Clean", "Start Year"),
    ("Last Updated (Raw)", "Last Updated"),
    ("Last Updated Year Clean", "Last Updated Year"),
    ("Update Recency", "Update Recency"),
    ("Start Decade", "Start Decade"),
    ("Modality", "Modality"),
    ("Project Home Country/Territory", "Project Home Country/Territory"),
    ("Location (Raw)", "Location"),
    ("City", "City"),
    ("State/Region", "State/Region"),
    ("Country", "Country"),
    ("Latitude", "Latitude"),
    ("Longitude", "Longitude"),
    ("Coordinate Stack Size", "Coordinate Stack Size"),
    ("Shared Coordinate?", "Shared Coordinate"),
    ("Named Contributor Band", "Named Contributor Band"),
    ("Project Type Count", "Project Type Count"),
    ("Topic Count", "Topic Count"),
    ("Era Count", "Era Count"),
    ("Organization Type Count", "Organization Type Count"),
    ("Contributor Type Count", "Contributor Type Count"),
    ("Leadership Tag Count", "Leadership Tag Count"),
    ("Funding Type Count", "Funding Type Count"),
    ("BIPOC Run?", "BIPOC Run"),
    ("Women Run?", "Women Run"),
    ("Leadership Data Available?", "Leadership Data Available"),
    ("Temporal Data Status", "Temporal Data Status"),
    ("Location Detail Status", "Location Detail Status"),
]

TAG_SPECS = [
    ("Project Type", "Project Type", "Project Types"),
    ("Topics Covered", "Topics Covered", "Topics"),
    ("Era Studied", "Era Studied", "Eras Studied"),
    ("Organization Type", "Organization Type", "Organization Types"),
    ("Creators and Contributors", "Creators and Contributors", "Contributor Roles"),
    ("Minority Leadership", "Minority Leadership", "Leadership Classifications"),
    ("FundingResources", "Funding/Resources", "Funding Resources"),
]


def read_rows(worksheet):
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).split())


def clean_year(value):
    """Convert cached formula zeros for missing years into empty CSV cells."""
    if value in {None, 0, "0"}:
        return ""
    return clean(value)


def tags_by_project(workbook):
    tags = {}
    for sheet_name, column_name, output_name in TAG_SPECS:
        values = defaultdict(set)
        for row in read_rows(workbook[sheet_name]):
            project_id = clean(row.get("BDPH ID"))
            value = clean(row.get(column_name))
            if project_id and value:
                values[project_id].add(value)
        tags[output_name] = values
    return tags


def export_projects(workbook):
    projects = read_rows(workbook["Visualization Ready"])
    tags = tags_by_project(workbook)
    columns = [output_name for _, output_name in PROJECT_FIELDS] + list(tags)

    with PROJECT_OUTPUT.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for project in projects:
            project_id = clean(project.get("BDPH ID"))
            row = {}
            for source_name, output_name in PROJECT_FIELDS:
                value = project.get(source_name)
                row[output_name] = (
                    clean_year(value)
                    if output_name in {"Start Year", "Last Updated Year"}
                    else clean(value)
                )
            row.update(
                {
                    output_name: "; ".join(sorted(values.get(project_id, set())))
                    for output_name, values in tags.items()
                }
            )
            writer.writerow(row)
    return len(projects)


def export_aggregates():
    prefix = "export const mappingVisualizations = "
    source = AGGREGATE_SOURCE.read_text(encoding="utf-8").strip()
    if not source.startswith(prefix) or not source.endswith(";"):
        raise ValueError("The visualization aggregate module has an unexpected format.")
    payload = json.loads(source[len(prefix) : -1])
    AGGREGATE_OUTPUT.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(f"Source workbook not found: {SOURCE}")
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    project_count = export_projects(workbook)
    export_aggregates()
    print(
        json.dumps(
            {
                "projects": project_count,
                "projectCsv": str(PROJECT_OUTPUT),
                "aggregateJson": str(AGGREGATE_OUTPUT),
            }
        )
    )


if __name__ == "__main__":
    main()
