"""Build compact, browser-ready aggregates from the analysis-ready workbook."""

from __future__ import annotations

import json
import math
import urllib.request
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SITE_ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "outputs"
    / "019f8b8b-1bb8-73c2-a6d9-ab5b8ac2baea"
    / "Mapping_v2_analysis_ready.xlsx"
)
OUTPUT = SITE_ROOT / "data" / "mapping-visualizations.js"
WORLD_SOURCE = SITE_ROOT / "data" / "world-110m.json"
WORLD_URL = "https://cdn.jsdelivr.net/npm/world-atlas@2/countries-110m.json"


def records(workbook, sheet_name):
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


def normalize_number(value, digits=6):
    if not isinstance(value, (int, float)) or not math.isfinite(value):
        return None
    return round(float(value), digits)


def sorted_counts(values):
    return [
        {"name": name, "count": count}
        for name, count in Counter(values).most_common()
    ]


def main():
    workbook = load_workbook(SOURCE, data_only=True, read_only=False)
    if WORLD_SOURCE.exists():
        world_topology = json.loads(WORLD_SOURCE.read_text(encoding="utf-8"))
    else:
        with urllib.request.urlopen(WORLD_URL) as response:
            world_topology = json.load(response)
    projects = records(workbook, "Visualization Ready")
    project_ids = {row["BDPH ID"] for row in projects}
    project_by_id = {row["BDPH ID"]: row for row in projects}

    normalized_specs = {
        "types": ("Project Type", "Project Type"),
        "topics": ("Topics Covered", "Topics Covered"),
        "eras": ("Era Studied", "Era Studied"),
        "organizations": ("Organization Type", "Organization Type"),
        "contributors": ("Creators and Contributors", "Creators and Contributors"),
        "leadership": ("Minority Leadership", "Minority Leadership"),
        "funding": ("FundingResources", "Funding/Resources"),
    }
    tags = {}
    for key, (sheet_name, column_name) in normalized_specs.items():
        by_project = defaultdict(set)
        for row in records(workbook, sheet_name):
            project_id = row["BDPH ID"]
            value = row[column_name]
            if project_id in project_ids and value:
                by_project[project_id].add(str(value).strip())
        tags[key] = by_project

    # Clustered map.
    coordinate_groups = defaultdict(list)
    for project in projects:
        latitude = normalize_number(project["Latitude"])
        longitude = normalize_number(project["Longitude"])
        if latitude is None or longitude is None:
            continue
        coordinate_groups[(latitude, longitude)].append(project)

    map_groups = []
    for (latitude, longitude), grouped_projects in coordinate_groups.items():
        grouped_projects.sort(key=lambda item: item["Name"] or "")
        location_bits = []
        for field in ("City", "State/Region", "Country"):
            values = [
                str(item[field]).strip()
                for item in grouped_projects
                if item[field] and str(item[field]).strip() not in {"N/A", "Unknown"}
            ]
            if values:
                location_bits.append(Counter(values).most_common(1)[0][0])
        map_groups.append(
            {
                "lat": latitude,
                "lon": longitude,
                "count": len(grouped_projects),
                "location": ", ".join(location_bits) or "Location detail unavailable",
                "projects": [
                    {
                        "id": item["BDPH ID"],
                        "name": item["Name"],
                        "url": item["Project URL"],
                    }
                    for item in grouped_projects
                ],
            }
        )
    map_groups.sort(key=lambda item: (-item["count"], item["lat"], item["lon"]))
    shared_coordinate_projects = sum(
        item["count"] for item in map_groups if item["count"] > 1
    )

    # Project type by topic heatmap.
    type_counts = Counter()
    topic_counts = Counter()
    type_topic_counts = Counter()
    for project_id in project_ids:
        project_types = tags["types"][project_id]
        project_topics = tags["topics"][project_id]
        type_counts.update(project_types)
        topic_counts.update(project_topics)
        for project_type in project_types:
            for topic in project_topics:
                type_topic_counts[(project_type, topic)] += 1
    type_order = [name for name, _ in type_counts.most_common()]
    topic_order = [name for name, _ in topic_counts.most_common()]
    type_topic_matrix = [
        [type_topic_counts[(project_type, topic)] for topic in topic_order]
        for project_type in type_order
    ]

    # Topic co-occurrence matrix.
    topic_pair_counts = Counter()
    for project_id in project_ids:
        project_topics = sorted(tags["topics"][project_id])
        for topic_a, topic_b in combinations(project_topics, 2):
            topic_pair_counts[(topic_a, topic_b)] += 1
    topic_matrix = []
    for topic_a in topic_order:
        row = []
        for topic_b in topic_order:
            if topic_a == topic_b:
                row.append(topic_counts[topic_a])
            else:
                key = tuple(sorted((topic_a, topic_b)))
                row.append(topic_pair_counts[key])
        topic_matrix.append(row)

    # Organization by contributor ecosystem.
    organization_counts = Counter()
    contributor_counts = Counter()
    organization_contributor_counts = Counter()
    for project_id in project_ids:
        project_organizations = tags["organizations"][project_id]
        project_contributors = tags["contributors"][project_id]
        organization_counts.update(project_organizations)
        contributor_counts.update(project_contributors)
        for organization in project_organizations:
            for contributor in project_contributors:
                organization_contributor_counts[(organization, contributor)] += 1
    organization_order = [name for name, _ in organization_counts.most_common()]
    contributor_order = [name for name, _ in contributor_counts.most_common()]
    organization_contributor_matrix = [
        [
            organization_contributor_counts[(organization, contributor)]
            for contributor in contributor_order
        ]
        for organization in organization_order
    ]

    # Historical era by start decade with modality encoded in each cell.
    era_order = [
        "14th Century and Earlier",
        "15th Century",
        "16th Century",
        "17th Century",
        "18th Century",
        "19th Century",
        "20th Century",
        "21st Century",
        "Other/Unknown",
    ]
    decade_order = [
        "1940s",
        "1960s",
        "1970s",
        "1980s",
        "1990s",
        "2000s",
        "2010s",
        "2020s",
        "Unknown",
    ]
    modality_order = ["Digital", "Mixed", "Physical"]
    era_decade_modality_counts = Counter()
    for project_id in project_ids:
        project = project_by_id[project_id]
        decade = project["Start Decade"] or "Unknown"
        modality = project["Modality"] or "Unknown"
        for era in tags["eras"][project_id]:
            if era in era_order and decade in decade_order and modality in modality_order:
                era_decade_modality_counts[(era, decade, modality)] += 1
    era_decade_matrix = [
        [
            [
                era_decade_modality_counts[(era, decade, modality)]
                for modality in modality_order
            ]
            for decade in decade_order
        ]
        for era in era_order
    ]

    # Funding and leadership transparency.
    leadership_available = {
        project["BDPH ID"]: project["Leadership Data Available?"] == "Yes"
        for project in projects
    }
    leadership_counts = Counter()
    funding_counts = Counter()
    funding_by_leadership = Counter()
    for project_id in project_ids:
        leadership_counts.update(tags["leadership"][project_id])
        funding_counts.update(tags["funding"][project_id])
        leadership_status = (
            "Known" if leadership_available[project_id] else "Unknown"
        )
        for funding_type in tags["funding"][project_id]:
            funding_by_leadership[(funding_type, leadership_status)] += 1
    funding_order = [name for name, _ in funding_counts.most_common()]
    funding_rows = [
        {
            "name": funding_type,
            "total": funding_counts[funding_type],
            "knownLeadership": funding_by_leadership[(funding_type, "Known")],
            "unknownLeadership": funding_by_leadership[
                (funding_type, "Unknown")
            ],
        }
        for funding_type in funding_order
    ]
    leadership_known = sum(leadership_available.values())
    leadership_unknown = len(projects) - leadership_known

    payload = {
        "meta": {
            "source": "Mapping_v2_analysis_ready.xlsx",
            "projects": len(projects),
            "mapCoordinateGroups": len(map_groups),
            "mapSharedProjects": shared_coordinate_projects,
            "mapSharedGroups": sum(1 for item in map_groups if item["count"] > 1),
            "mapLargestStack": max(item["count"] for item in map_groups),
            "leadershipKnown": leadership_known,
            "leadershipUnknown": leadership_unknown,
            "unknownStartDecade": sum(
                1
                for project in projects
                if project["Start Decade"] in {None, "", "Unknown"}
            ),
        },
        "map": {"world": world_topology, "groups": map_groups},
        "typeTopic": {
            "types": [
                {"name": name, "count": type_counts[name]} for name in type_order
            ],
            "topics": [
                {"name": name, "count": topic_counts[name]} for name in topic_order
            ],
            "matrix": type_topic_matrix,
        },
        "topicCooccurrence": {
            "topics": [
                {"name": name, "count": topic_counts[name]} for name in topic_order
            ],
            "matrix": topic_matrix,
        },
        "organizationContributor": {
            "organizations": [
                {"name": name, "count": organization_counts[name]}
                for name in organization_order
            ],
            "contributors": [
                {"name": name, "count": contributor_counts[name]}
                for name in contributor_order
            ],
            "matrix": organization_contributor_matrix,
        },
        "eraModalityDecade": {
            "eras": era_order,
            "decades": decade_order,
            "modalities": modality_order,
            "matrix": era_decade_matrix,
        },
        "fundingLeadership": {
            "leadershipKnown": leadership_known,
            "leadershipUnknown": leadership_unknown,
            "leadershipTags": sorted_counts(
                tag
                for project_id in project_ids
                for tag in tags["leadership"][project_id]
            ),
            "funding": funding_rows,
        },
    }

    output = (
        "export const mappingVisualizations = "
        + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        + ";\n"
    )
    OUTPUT.write_text(output, encoding="utf-8")
    print(
        json.dumps(
            {
                "output": str(OUTPUT),
                "bytes": len(output.encode("utf-8")),
                "projects": len(projects),
                "sharedCoordinateProjects": shared_coordinate_projects,
                "largestStack": payload["meta"]["mapLargestStack"],
                "unknownLeadership": leadership_unknown,
            }
        )
    )


if __name__ == "__main__":
    main()
